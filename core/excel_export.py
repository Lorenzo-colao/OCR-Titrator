"""
Export Excel avec mise en forme et flags.

Génère un fichier ``.xlsx`` contenant :
- Feuille « Données » : tableau coloré selon les flags
- Feuille « ⚠ Flags à vérifier » : résumé des anomalies

Ce module ne dépend que de ``config``, ``core.validation``, ``openpyxl`` et ``pandas``.
Aucune référence à l'UI.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import CFG, COLONNES
from core.validation import (
    check_and_fix_cumulative,
    check_physical_consistency,
    compute_iqr_bounds,
    flag_row,
)

# ──────────────────────────────────────────────────────────────────────────────
#  Styles Excel
# ──────────────────────────────────────────────────────────────────────────────

_HEADER_FILL  = PatternFill("solid", fgColor="0F3460")
_HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
_RED_FILL     = PatternFill("solid", fgColor="C0392B")
_ORANGE_FILL  = PatternFill("solid", fgColor="E67E22")
_YELLOW_FILL  = PatternFill("solid", fgColor="F39C12")
_BLUE_FILL    = PatternFill("solid", fgColor="2471A3")
_PURPLE_FILL  = PatternFill("solid", fgColor="7D3C98")
_GREEN_FILL   = PatternFill("solid", fgColor="1E8449")
_RED_FONT     = Font(color="FFFFFF", bold=True, size=10)
_WHITE_FONT   = Font(color="FFFFFF", bold=True, size=10)
_BLUE_FONT    = Font(color="FFFFFF", bold=True, size=10)
_PURPLE_FONT  = Font(color="FFFFFF", bold=True, size=10)
_OK_FONT      = Font(color="AAAAAA", italic=True, size=9)
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_ALIGN   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

_NUM_FORMATS: dict[str, str] = {
    "mL":    "0.0000",
    "mL2":   "0.0000",
    "pH":    "0.000",
    "delta": "0.000",
    "S":     "0.0000",
    "C":     "0.0",
}

_COL_WIDTHS: dict[str, int] = {
    "mL": 11, "mL2": 11, "pH": 9, "delta": 11,
    "S": 9, "C": 9, "Flag": 18,
}


# ──────────────────────────────────────────────────────────────────────────────
#  Export principal
# ──────────────────────────────────────────────────────────────────────────────

def validate_dataframe_for_export(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, tuple[float, float]], int]:
    """Applique toutes les validations et retourne le DataFrame flaggé, les stats et le nombre d'anomalies."""
    df = df.copy()
    df = check_and_fix_cumulative(df)
    df = check_physical_consistency(df)
    col_stats = compute_iqr_bounds(df)
    df["Flag"] = df.apply(lambda r: flag_row(r, col_stats), axis=1)

    # Fusionner Phys_Flag dans Flag
    for i in df.index:
        phys = df.at[i, "Phys_Flag"]
        if phys:
            sep = " | " if df.at[i, "Flag"] else ""
            df.at[i, "Flag"] = df.at[i, "Flag"] + sep + phys

    # Fusionner corrections dans Flag
    for i in df.index:
        fixes = [df.at[i, c] for c in ("Fix_mL2", "Fix_delta") if df.at[i, c]]
        if fixes:
            sep = " | " if df.at[i, "Flag"] else ""
            df.at[i, "Flag"] = df.at[i, "Flag"] + sep + " | ".join(fixes)

    df_reset = df.reset_index(drop=True)
    n_flags = len(df_reset[
        (df_reset["Flag"] != "")
        | (df_reset["Fix_mL2"] != "")
        | (df_reset["Fix_delta"] != "")
        | (df_reset["Phys_Flag"] != "")
    ])
    return df, col_stats, n_flags


def export_excel_with_flags(df: pd.DataFrame, out_path: str) -> int:
    """
    Exporte un DataFrame en Excel avec validation complète et mise en forme.

    Args:
        df: DataFrame brut (colonnes ``COLONNES``).
        out_path: Chemin du fichier ``.xlsx`` de sortie.

    Returns:
        Nombre de lignes flaggées (anomalies + corrections).
    """
    df, col_stats, n_flags = validate_dataframe_for_export(df)

    wb = Workbook()
    _write_data_sheet(wb, df, col_stats)
    _write_flags_sheet(wb, df)
    wb.save(out_path)

    return n_flags


# ──────────────────────────────────────────────────────────────────────────────
#  Feuille « Données »
# ──────────────────────────────────────────────────────────────────────────────

def _write_data_sheet(
    wb: Workbook,
    df: pd.DataFrame,
    col_stats: dict[str, tuple[float, float]],
) -> None:
    """Écrit la feuille principale avec les données et le formatage conditionnel."""
    ws = wb.active
    ws.title = "Données"
    ws.row_dimensions[1].height = 22

    _ALIGN_NO_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=False)

    all_cols = COLONNES

    # En-tête
    for j, col in enumerate(all_cols, 1):
        cell = ws.cell(row=1, column=j, value=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _ALIGN_NO_WRAP

    # Corps
    for excel_row, (_, row_data) in enumerate(df.iterrows(), start=2):
        flag_msg   = row_data.get("Flag", "")
        fix_ml2    = row_data.get("Fix_mL2",   "")
        fix_delta  = row_data.get("Fix_delta", "")
        has_flag   = bool(flag_msg)
        is_fix_row = has_flag and "corrigé" in flag_msg

        for j, col in enumerate(all_cols, 1):
            val  = row_data[col]
            cell = ws.cell(row=excel_row, column=j)
            cell.alignment = _ALIGN_NO_WRAP

            _format_data_cell(
                cell, val, col, col_stats,
                has_flag, is_fix_row, fix_ml2, fix_delta,
            )

    # Largeurs de colonnes
    for j, col in enumerate(all_cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = _COL_WIDTHS.get(col, 12)
    ws.freeze_panes = "A2"


def _format_data_cell(
    cell,
    val,
    col: str,
    col_stats: dict[str, tuple[float, float]],
    has_flag: bool,
    is_fix_row: bool,
    fix_ml2: str,
    fix_delta: str,
) -> None:
    """Applique le style conditionnel à une cellule de données."""
    if pd.isna(val):
        cell.value = None
        if has_flag and not is_fix_row:
            cell.fill = _ORANGE_FILL
    else:
        cell.value         = float(val)
        cell.number_format = _NUM_FORMATS.get(col, "0.0000")
        if col == "mL2" and fix_ml2:
            cell.fill, cell.font = _BLUE_FILL, _BLUE_FONT
        elif col == "delta" and fix_delta:
            cell.fill, cell.font = _BLUE_FILL, _BLUE_FONT
        elif has_flag and not is_fix_row:
            v_rules = CFG.validation_rules.get(col)
            if v_rules and (val < v_rules["min"] or val > v_rules["max"]):
                cell.fill, cell.font = _RED_FILL, _WHITE_FONT
            elif col_stats and col in col_stats:
                lo, hi = col_stats[col]
                if val < lo or val > hi:
                    cell.fill, cell.font = _YELLOW_FILL, _WHITE_FONT


# ──────────────────────────────────────────────────────────────────────────────
#  Feuille « ⚠ Flags à vérifier »
# ──────────────────────────────────────────────────────────────────────────────

def _write_flags_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    """Écrit la feuille résumé des anomalies."""
    ws2 = wb.create_sheet("⚠ Flags à vérifier")
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 75

    for j, title in enumerate(["Ligne Excel", "Type", "Détail"], 1):
        c = ws2.cell(1, j, title)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = _CENTER_ALIGN

    df_reset = df.reset_index(drop=True)
    k = 2
    for pos in range(len(df_reset)):
        row_data       = df_reset.iloc[pos]
        flag_msg       = row_data["Flag"]
        fix_ml2        = row_data.get("Fix_mL2",   "")
        fix_delta      = row_data.get("Fix_delta", "")
        phys_msg       = row_data.get("Phys_Flag", "")
        excel_data_row = pos + 2

        # Corrections
        for fix_msg, fix_col in [(fix_ml2, "mL2 corrigé"), (fix_delta, "delta corrigé")]:
            if fix_msg:
                ws2.cell(k, 1, excel_data_row).fill = _BLUE_FILL
                ws2.cell(k, 1).font = _WHITE_FONT
                ws2.cell(k, 1).alignment = _CENTER_ALIGN
                ws2.cell(k, 2, fix_col).fill = _BLUE_FILL
                ws2.cell(k, 2).font = _BLUE_FONT
                ws2.cell(k, 2).alignment = _CENTER_ALIGN
                ws2.cell(k, 3, fix_msg).alignment = _LEFT_ALIGN
                k += 1

        # Flags physiques
        if phys_msg:
            for rule_detail in phys_msg.split(" | "):
                rule_detail = rule_detail.strip()
                if not rule_detail:
                    continue
                ws2.cell(k, 1, excel_data_row).fill = _PURPLE_FILL
                ws2.cell(k, 1).font = _WHITE_FONT
                ws2.cell(k, 1).alignment = _CENTER_ALIGN
                ws2.cell(k, 2, "Cohérence phys.").fill = _PURPLE_FILL
                ws2.cell(k, 2).font = _PURPLE_FONT
                ws2.cell(k, 2).alignment = _CENTER_ALIGN
                ws2.cell(k, 3, rule_detail).alignment = _LEFT_ALIGN
                k += 1

        # Flags de données
        flag_display = flag_msg
        for fix in [fix_ml2, fix_delta]:
            if fix:
                flag_display = flag_display.replace(" | " + fix, "").replace(fix, "").strip(" |")
        if phys_msg:
            for part in phys_msg.split(" | "):
                flag_display = flag_display.replace(
                    " | " + part.strip(), ""
                ).replace(part.strip(), "")
        flag_display = flag_display.strip(" |")

        if flag_display:
            if "manquant" in flag_display:
                typ, fill = "Valeur NaN", _ORANGE_FILL
            elif "aberrante" in flag_display:
                typ, fill = "Outlier stat.", _YELLOW_FILL
            else:
                typ, fill = "Hors plage", _RED_FILL
            ws2.cell(k, 1, excel_data_row).fill = fill
            ws2.cell(k, 1).font = _WHITE_FONT
            ws2.cell(k, 1).alignment = _CENTER_ALIGN
            ws2.cell(k, 2, typ).fill = fill
            ws2.cell(k, 2).font = _WHITE_FONT
            ws2.cell(k, 2).alignment = _CENTER_ALIGN
            ws2.cell(k, 3, flag_display).alignment = _LEFT_ALIGN
            k += 1
